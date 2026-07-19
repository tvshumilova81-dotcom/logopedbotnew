from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from database.models.booking import Booking
from database.models.material import Material
from database.models.purchase import Purchase
from database.models.questionnaire import Questionnaire
from database.models.user import User


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


async def export_all(session: AsyncSession, output_path: Path) -> Path:
    wb = Workbook()

    # Пользователи
    ws_users = wb.active
    ws_users.title = "Пользователи"
    _write_header(
        ws_users,
        ["ID", "Telegram ID", "Username", "Родитель", "Ребёнок", "Возраст",
         "Страна", "Часовой пояс", "Дата регистрации"],
    )
    users = (await session.execute(select(User))).scalars().all()
    users_by_id = {u.id: u for u in users}
    for u in users:
        ws_users.append([
            u.id, u.telegram_id, u.username or "", u.parent_name or "",
            u.child_name or "", u.child_age or "", u.country or "",
            u.timezone or "", u.created_at.strftime("%d.%m.%Y %H:%M"),
        ])

    # Заявки
    ws_bookings = wb.create_sheet("Заявки")
    _write_header(
        ws_bookings,
        ["ID", "Родитель", "Ребёнок", "Возраст", "Страна", "Проблема",
         "Статус", "Дата"],
    )
    bookings = (await session.execute(select(Booking))).scalars().all()
    for b in bookings:
        ws_bookings.append([
            b.id, b.parent_name, b.child_name, b.child_age, b.country,
            b.problem or "", b.status.value, b.created_at.strftime("%d.%m.%Y %H:%M"),
        ])

    # Анкеты
    ws_q = wb.create_sheet("Анкеты")
    _write_header(
        ws_q,
        ["ID", "Пользователь", "Ребёнок", "Возраст", "Заполнена", "Дата"],
    )
    questionnaires = (await session.execute(select(Questionnaire))).scalars().all()
    for q in questionnaires:
        owner = users_by_id.get(q.user_id)
        ws_q.append([
            q.id, owner.display_name() if owner else "", q.child_name or "",
            q.child_age or "", "Да" if q.is_completed else "Нет",
            q.created_at.strftime("%d.%m.%Y %H:%M"),
        ])

    # Покупки
    ws_p = wb.create_sheet("Покупки")
    _write_header(ws_p, ["ID", "Пользователь", "Материал", "Stars", "Дата"])
    purchases = (await session.execute(select(Purchase))).scalars().all()
    materials = {m.id: m for m in (await session.execute(select(Material))).scalars().all()}
    for p in purchases:
        owner = users_by_id.get(p.user_id)
        material = materials.get(p.material_id)
        ws_p.append([
            p.id, owner.display_name() if owner else "",
            material.title if material else "", p.price_stars,
            p.created_at.strftime("%d.%m.%Y %H:%M"),
        ])

    wb.save(output_path)
    return output_path
